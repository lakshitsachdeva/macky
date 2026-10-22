from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable


FIRST_PAGE_LENGTH = 500
SUMMARY_PAGE_LENGTH = 1000
MAX_HEADINGS = 20
SUMMARY_HEADINGS = 10
MAX_ROWS_PER_SHEET = 1000
MAX_TOTAL_CHARS = 50_000
MAX_CELLS_PER_ROW = 20
BLANK_LINE_PATTERN = re.compile(r"(?:\n\s*){3,}")
MARKDOWN_HEADING_PATTERN = re.compile(r"^\s{0,3}#{1,6}\s+(.*\S)\s*$")


@dataclass
class ExtractedContent:
    full_text: str = ""
    title: str = ""
    headings: list[str] = field(default_factory=list)
    first_page: str = ""
    summary_text: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)
    error: str | None = None


def _clean_text(text: str) -> str:
    normalized = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    normalized = BLANK_LINE_PATTERN.sub("\n\n", normalized)
    return normalized.strip()


def _dedupe_preserve(values: list[str]) -> list[str]:
    deduped: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = _clean_text(value)
        if not cleaned:
            continue
        key = cleaned.casefold()
        if key in seen:
            continue
        seen.add(key)
        deduped.append(cleaned)
    return deduped


def _empty_content(file_path: Path, error: str | None = None) -> ExtractedContent:
    return ExtractedContent(
        full_text="",
        title=file_path.stem,
        headings=[],
        first_page="",
        summary_text=file_path.stem,
        metadata={},
        error=error,
    )


def _finalize_content(
    file_path: Path,
    *,
    full_text: str,
    title: str,
    headings: list[str] | None = None,
    first_page: str,
    summary_parts: list[str] | None = None,
    metadata: dict[str, Any] | None = None,
    error: str | None = None,
) -> ExtractedContent:
    cleaned_full_text = _clean_text(full_text)
    cleaned_title = _clean_text(title) or file_path.stem
    cleaned_headings = _dedupe_preserve((headings or [])[:MAX_HEADINGS])
    cleaned_first_page = _clean_text(first_page) or cleaned_full_text[:FIRST_PAGE_LENGTH]

    normalized_summary_parts = [_clean_text(part) for part in (summary_parts or []) if part]
    summary_text = _clean_text(" ".join(part for part in normalized_summary_parts if part))
    if not summary_text:
        summary_text = _clean_text(
            " ".join(
                filter(
                    None,
                    [
                        cleaned_title,
                        " ".join(cleaned_headings[:SUMMARY_HEADINGS]),
                        cleaned_first_page[:SUMMARY_PAGE_LENGTH],
                    ],
                )
            )
        )
    if not summary_text:
        summary_text = cleaned_full_text[:SUMMARY_PAGE_LENGTH] or cleaned_title

    return ExtractedContent(
        full_text=cleaned_full_text,
        title=cleaned_title,
        headings=cleaned_headings,
        first_page=cleaned_first_page[:FIRST_PAGE_LENGTH],
        summary_text=summary_text,
        metadata=metadata or {},
        error=error,
    )


def _extract_pdf_structured(file_path: Path) -> ExtractedContent:
    import fitz

    with fitz.open(str(file_path)) as document:
        metadata = document.metadata or {}
        title = (metadata.get("title") or "").strip()
        if title.casefold() in {"untitled", "unknown", "document"}:
            title = ""
        headings: list[str] = []
        page0_text_parts: list[str] = []
        max_font_size = 0.0

        if document.page_count > 0:
            first_page = document[0]
            blocks = first_page.get_text("dict").get("blocks", [])
            for block in blocks:
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    line_parts: list[str] = []
                    line_max_font = 0.0
                    for span in line.get("spans", []):
                        span_text = (span.get("text") or "").strip()
                        if not span_text:
                            continue
                        font_size = float(span.get("size") or 0.0)
                        line_parts.append(span_text)
                        page0_text_parts.append(span_text)
                        line_max_font = max(line_max_font, font_size)
                        if font_size > max_font_size and len(span_text) > 3 and not title:
                            max_font_size = font_size
                            title = span_text

                    line_text = " ".join(line_parts).strip()
                    if line_text and line_max_font >= 12 and len(line_text) < 100:
                        headings.append(line_text)

        full_text = "\n".join(page.get_text() for page in document)
        page0_text = " ".join(page0_text_parts)

        return _finalize_content(
            file_path,
            full_text=full_text,
            title=title or file_path.stem,
            headings=headings,
            first_page=page0_text,
            summary_parts=[
                title,
                " ".join(headings[:SUMMARY_HEADINGS]),
                page0_text[:SUMMARY_PAGE_LENGTH],
            ],
            metadata={
                "page_count": document.page_count,
                "author": metadata.get("author", ""),
                "subject": metadata.get("subject", ""),
            },
        )


def _extract_docx_structured(file_path: Path) -> ExtractedContent:
    from docx import Document

    document = Document(str(file_path))
    headings: list[str] = []
    body_text: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
        if style_name.startswith("Heading"):
            headings.append(text)
        else:
            body_text.append(text)

    table_text = [
        cell.text.strip()
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    ]

    core_title = (document.core_properties.title or "").strip()
    title = core_title or (headings[0] if headings else "")
    if not title:
        title = body_text[0] if body_text and len(body_text[0]) < 120 else file_path.stem

    full_text = "\n".join([*headings, *body_text, *table_text])
    first_page = "\n".join(body_text[:5] or headings[:5] or table_text[:5])

    return _finalize_content(
        file_path,
        full_text=full_text,
        title=title,
        headings=headings,
        first_page=first_page,
        summary_parts=[
            title,
            " ".join(headings[:SUMMARY_HEADINGS]),
            "\n".join(body_text[:3] or table_text[:3]),
        ],
        metadata={
            "paragraph_count": len(body_text),
            "heading_count": len(headings),
            "table_cell_count": len(table_text),
        },
    )


def _extract_plain_text_structured(file_path: Path) -> ExtractedContent:
    encoding = "utf-8"
    try:
        text = file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        encoding = "latin-1"
        text = file_path.read_text(encoding="latin-1")

    lines = [line.rstrip() for line in text.splitlines()]
    non_empty_lines = [line.strip() for line in lines if line.strip()]
    headings = [
        match.group(1).strip()
        for match in (MARKDOWN_HEADING_PATTERN.match(line) for line in lines)
        if match
    ]
    title_candidate = headings[0] if headings else (non_empty_lines[0] if non_empty_lines else "")
    title = title_candidate if title_candidate and len(title_candidate) < 120 else file_path.stem
    if not headings:
        headings = [
            line
            for line in non_empty_lines[:20]
            if len(line) < 100 and (line.endswith(":") or line.isupper())
        ]

    first_page = "\n".join(non_empty_lines[:5])
    return _finalize_content(
        file_path,
        full_text=text,
        title=title,
        headings=headings,
        first_page=first_page,
        summary_parts=[
            title,
            " ".join(headings[:SUMMARY_HEADINGS]),
            first_page[:SUMMARY_PAGE_LENGTH],
        ],
        metadata={"encoding": encoding},
    )


def _extract_xlsx_structured(file_path: Path) -> ExtractedContent:
    import openpyxl

    workbook = openpyxl.load_workbook(
        str(file_path),
        read_only=True,
        data_only=True,
    )
    try:
        sheet_texts: list[str] = []
        sheet_names = [sheet.title for sheet in workbook.worksheets]
        total_chars = 0

        for sheet in workbook.worksheets:
            if total_chars >= MAX_TOTAL_CHARS:
                break

            sheet_lines = [f"Sheet: {sheet.title}"]
            rows_read = 0

            for row in sheet.iter_rows(
                max_row=MAX_ROWS_PER_SHEET,
                values_only=True,
            ):
                if rows_read >= MAX_ROWS_PER_SHEET or total_chars >= MAX_TOTAL_CHARS:
                    break
                if not row:
                    continue

                cells = [
                    str(cell)
                    for cell in row
                    if cell is not None and str(cell).strip()
                ]
                if not cells:
                    continue

                if len(cells) > MAX_CELLS_PER_ROW:
                    cells = cells[:MAX_CELLS_PER_ROW]

                row_text = " | ".join(cells)
                sheet_lines.append(row_text)
                total_chars += len(row_text)
                rows_read += 1

            sheet_texts.append("\n".join(sheet_lines))

        full_text = "\n\n".join(sheet_texts)
        title = sheet_names[0] if sheet_names else file_path.stem
        first_page = sheet_texts[0] if sheet_texts else ""
        return _finalize_content(
            file_path,
            full_text=full_text,
            title=title,
            headings=sheet_names,
            first_page=first_page,
            summary_parts=[
                title,
                " ".join(sheet_names[:SUMMARY_HEADINGS]),
                first_page[:SUMMARY_PAGE_LENGTH],
            ],
            metadata={"sheet_names": sheet_names, "sheet_count": len(sheet_names)},
        )
    finally:
        workbook.close()


def _extract_pptx_structured(file_path: Path) -> ExtractedContent:
    from pptx import Presentation

    presentation = Presentation(str(file_path))
    slide_texts: list[str] = []
    headings: list[str] = []

    for index, slide in enumerate(presentation.slides):
        shapes_text = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        if not shapes_text:
            continue

        candidate_heading = shapes_text[0].splitlines()[0].strip()
        if candidate_heading and len(candidate_heading) < 120:
            headings.append(candidate_heading)

        joined_text = " | ".join(shapes_text)
        slide_texts.append(f"Slide {index + 1}: {joined_text}")

    full_text = "\n".join(slide_texts)
    title = headings[0] if headings else file_path.stem
    first_page = slide_texts[0] if slide_texts else ""
    return _finalize_content(
        file_path,
        full_text=full_text,
        title=title,
        headings=headings,
        first_page=first_page,
        summary_parts=[
            title,
            " ".join(headings[:SUMMARY_HEADINGS]),
            first_page[:SUMMARY_PAGE_LENGTH],
        ],
        metadata={"slide_count": len(presentation.slides)},
    )


HANDLERS: dict[str, Callable[[Path], ExtractedContent]] = {
    ".pdf": _extract_pdf_structured,
    ".docx": _extract_docx_structured,
    ".doc": _extract_docx_structured,
    ".txt": _extract_plain_text_structured,
    ".md": _extract_plain_text_structured,
    ".csv": _extract_plain_text_structured,
    ".xlsx": _extract_xlsx_structured,
    ".pptx": _extract_pptx_structured,
}


def extract_structured(file_path: Path) -> ExtractedContent:
    file_path = Path(file_path)
    handler = HANDLERS.get(file_path.suffix.lower())
    if handler is None:
        return _empty_content(
            file_path,
            error=f"Unsupported file type: {file_path.suffix or 'no extension'}",
        )

    try:
        return handler(file_path)
    except Exception as exc:
        return _empty_content(file_path, error=str(exc))


def extract_content(file_path: Path) -> dict[str, Any]:
    content = extract_structured(file_path)
    return {
        "text": content.full_text,
        "title": content.title,
        "metadata": content.metadata,
        "error": content.error,
        "preview_text": content.first_page,
        "headings": content.headings,
        "summary_text": content.summary_text,
    }
